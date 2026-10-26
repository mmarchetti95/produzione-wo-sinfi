#!/usr/bin/env python3
"""
Report produzione rilievi da QFieldCloud (EagleProjects) - motore di calcolo
============================================================================
Legge i "delta" (modifiche sincronizzate da QField) di tutti i progetti
visibili all'utente configurato e calcola, per giorno / operatore / progetto / layer:
  - km netti di linee   = nuove + allungamenti - accorciamenti - cancellazioni
  - km nuovi            = solo linee create (lordo)
  - feature create / modificate / cancellate, e quanti dei creati sono punti
  - tipo di geometria di ogni layer (punto, linea, poligono)
  - foto scattate, orario della prima (inizio) e dell'ultima (fine) foto della giornata
  - orario dell'ultima sincronizzazione (ultimo delta arrivato sul server)
(calcola internamente anche m2 di poligoni e delta in errore, oggi non mostrati)

Produce un Excel oppure (con --json) i dati della board web in sito/.
La GitHub Action del repository lo lancia periodicamente e pubblica la board su GitHub Pages.
I giorni sono sempre in ora italiana e sono quelli di SINCRONIZZAZIONE, tranne per foto/inizio/fine
che vanno nel giorno di SCATTO, letto dal nome del file (es. DCIM/JPEG_20261015083012345.jpg).

Uso:
  python qfc_report.py                          # ultimi 7 giorni
  python qfc_report.py --giorni 60              # ultimi 60 giorni
  python qfc_report.py --dal 2026-10-01 --al 2026-10-12
  python qfc_report.py --progetti "Comune_X" "Comune_Y"
  python qfc_report.py --tutti                  # non salta i progetti inattivi
  python qfc_report.py --giorni 60 --json sito/dati.json   # dati per la board web

Credenziali: variabili d'ambiente (Secrets della GitHub Action) o file .env nella stessa cartella
  QFC_URL=https://qfield.eagleprojects.cloud
  QFC_USER=utente_reporting
  QFC_PASSWORD=...
Dipendenze: requests, shapely, pyproj, openpyxl  (tzdata su Windows)
"""
import argparse
import json
import os
import re
import sys
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

try:
    import requests
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from pyproj import Geod
    from requests.adapters import HTTPAdapter
    from shapely import wkt as shp_wkt
    from urllib3.util.retry import Retry
except ImportError as e:
    sys.exit(f"\nManca la libreria Python '{e.name}' per questo Python:\n  {sys.executable}\n"
             f"Installa le librerie con:\n  \"{sys.executable}\" -m pip install requests shapely pyproj openpyxl tzdata\n"
             f"oppure avvia la board con avvia_board.bat, che le installa da solo.\n")

try:
    from zoneinfo import ZoneInfo
    TZ = ZoneInfo("Europe/Rome")
except Exception:  # Windows senza il pacchetto tzdata: uso il fuso del PC
    TZ = None

GEOD = Geod(ellps="WGS84")
# delta che consideriamo "lavoro arrivato sul server" (contano nei km)
STATI_VALIDI = {"applied", "conflict", "not_applied", "pending", "started"}
# delta da segnalare al coordinatore
STATI_PROBLEMA = {"conflict", "not_applied", "error", "unpermitted"}
KM_MAX_FEATURE = 50      # una feature più lunga è quasi certamente un errore di CRS
MARGINE_INATTIVI = 30    # giorni: progetti fermi da prima di (dal - margine) vengono saltati
WORKERS = 6              # download in parallelo
METRICHE = ("km_netti", "km_nuovi", "mq_netti", "create", "punti", "modificate", "cancellate", "problemi", "foto")
TIPI = {"punto": "Punto", "linea": "Linea", "poligono": "Poligono", "": "n.d."}  # geometria del layer
RE_FOTO = re.compile(r"\.(jpe?g|png|heic|heif|webp)$", re.I)
# data e ora nel nome del file: JPEG_20261015083012345.jpg, IMG_20261015_083012.jpg, 2026-10-15_08-30-12.jpg
RE_ORA_FOTO = re.compile(r"(20\d\d)[-_]?(\d\d)[-_]?(\d\d)[-_T ]?(\d\d)[-_.:]?(\d\d)[-_.:]?(\d\d)")


def locale(dt: datetime) -> datetime:
    return dt.astimezone(TZ) if TZ else dt.astimezone()


def adesso() -> datetime:
    return locale(datetime.now(timezone.utc))


def oggi() -> date:
    return adesso().date()


def ts(s):
    if not s:
        return None
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return None
    return d if d.tzinfo else d.replace(tzinfo=timezone.utc)


# ----------------------------------------------------------------- config ---
def carica_env():
    env = Path(__file__).with_name(".env")
    if env.exists():
        for riga in env.read_text(encoding="utf-8").splitlines():
            riga = riga.strip()
            if riga and not riga.startswith("#") and "=" in riga:
                k, v = riga.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    url = (os.environ.get("QFC_URL") or "https://qfield.eagleprojects.cloud").rstrip("/")
    user, pwd = os.environ.get("QFC_USER"), os.environ.get("QFC_PASSWORD")
    if not (user and pwd):
        raise RuntimeError("Imposta QFC_USER e QFC_PASSWORD (variabili d'ambiente o file .env)")
    return url, user, pwd


# -------------------------------------------------------------------- API ---
class QFC:
    def __init__(self, url, user, pwd):
        self.api = f"{url}/api/v1"
        self._tl = threading.local()
        r = self._s().post(f"{self.api}/auth/login/", json={"username": user, "password": pwd}, timeout=60)
        if r.status_code in (400, 401, 403):
            raise RuntimeError("Login QFieldCloud fallito: controlla utente e password nel file .env")
        r.raise_for_status()
        self.token = r.json()["token"]
        self._s().headers["Authorization"] = f"Token {self.token}"

    def _s(self):
        """Una sessione per thread, con retry automatico sugli errori temporanei."""
        s = getattr(self._tl, "s", None)
        if s is None:
            s = requests.Session()
            retry = Retry(total=4, backoff_factor=1.5, status_forcelist=(429, 500, 502, 503, 504),
                          allowed_methods=("GET", "POST"))
            s.mount("https://", HTTPAdapter(max_retries=retry))
            s.mount("http://", HTTPAdapter(max_retries=retry))
            if getattr(self, "token", None):
                s.headers["Authorization"] = f"Token {self.token}"
            self._tl.s = s
        return s

    def _get_tutti(self, path, params=None):
        """GET che gestisce sia liste semplici sia risposte paginate."""
        out, params = [], dict(params or {})
        url = f"{self.api}{path}"
        while url:
            r = self._s().get(url, params=params, timeout=300)
            r.raise_for_status()
            data = r.json()
            if isinstance(data, list):
                out.extend(data)
                break
            out.extend(data.get("results", []))
            url, params = data.get("next"), None
        return out

    def progetti(self):
        return self._get_tutti("/projects/")

    def delta(self, project_id):
        return self._get_tutti(f"/deltas/{project_id}/")


# -------------------------------------------------------------- geometria ---
def misura(geom_wkt):
    """(km_linee, m2_poligoni, scartata). Gestisce CRS metrici e geografici."""
    if not geom_wkt:
        return 0.0, 0.0, False
    try:
        g = shp_wkt.loads(geom_wkt)
    except Exception:
        return 0.0, 0.0, False
    if g.is_empty:
        return 0.0, 0.0, False
    minx, miny, maxx, maxy = g.bounds
    geografico = -180 <= minx <= 180 and -90 <= miny <= 90 and -180 <= maxx <= 180 and -90 <= maxy <= 90
    t = g.geom_type
    if "Line" in t:
        km = (GEOD.geometry_length(g) if geografico else g.length) / 1000.0
        if km > KM_MAX_FEATURE:
            return 0.0, 0.0, True
        return km, 0.0, False
    if "Polygon" in t:
        mq = abs(GEOD.geometry_area_perimeter(g)[0]) if geografico else g.area
        return 0.0, mq, False
    return 0.0, 0.0, False


def tipo_geom(geom_wkt):
    """'punto' | 'linea' | 'poligono' dalla prima parola del WKT, None se manca o non si riconosce."""
    m = re.match(r"\s*([A-Za-z]+)", geom_wkt or "")
    if not m:
        return None
    t = m.group(1).upper()
    if "POINT" in t:
        return "punto"
    if "POLYGON" in t or "SURFACE" in t:
        return "poligono"
    if "LINE" in t or "CURVE" in t:
        return "linea"
    return None


def foto_delta(c):
    """Percorsi delle foto allegate al delta (file inviati e attributi che puntano a un'immagine)."""
    new = c.get("new") or {}
    nomi = set((new.get("files_sha256") or {}).keys())
    nomi.update(v.strip() for v in (new.get("attributes") or {}).values() if isinstance(v, str))
    return {n for n in nomi if RE_FOTO.search(n)}


def nome_file(percorso):
    return percorso.replace("\\", "/").rsplit("/", 1)[-1]


def ora_foto(percorso, sincronizzata: datetime):
    """Data e ora di scatto (orologio del dispositivo) lette dal nome del file.
    None se il nome non le contiene o se non sono plausibili rispetto alla sincronizzazione."""
    m = RE_ORA_FOTO.search(nome_file(percorso))
    if not m:
        return None
    try:
        dt = datetime(*map(int, m.groups()))
    except ValueError:
        return None
    s = sincronizzata.replace(tzinfo=None)
    return dt if s - timedelta(days=60) <= dt <= s + timedelta(hours=2) else None


def ha_sync(r):
    """La riga contiene lavoro sincronizzato (non solo foto scattate quel giorno e sincronizzate dopo)."""
    return any(r[m] for m in METRICHE if m != "foto")


def ore_lavoro(r):
    """Orario lavorativo in ore: dalla prima all'ultima foto (None se non ci sono foto)."""
    if not r.get("inizio"):
        return None
    (h0, m0), (h1, m1) = (map(int, r[k].split(":")) for k in ("inizio", "fine"))
    return (h1 * 60 + m1 - h0 * 60 - m0) / 60


def ultima_attivita(p):
    """Data (locale) dell'ultima attività nota del progetto, None se il server non la espone."""
    date_ = [ts(p.get(k)) for k in ("data_last_updated_at", "updated_at", "data_last_packaged_at")]
    date_ = [d for d in date_ if d]
    return locale(max(date_)).date() if date_ else None


# ------------------------------------------------------------- elaborazione ---
def elabora(client, dal: date, al: date, filtro_progetti=None, tutti=False,
            log=print, avanzamento=None):
    """
    Ritorna un dict:
      righe     lista di dict (giorno, operatore, progetto, layer, geom + METRICHE + inizio/fine/sync "HH:MM")
      avvisi    lista di stringhe
      progetti  {"letti": n, "saltati": n, "errori": n}
    """
    progetti = client.progetti()
    if filtro_progetti:
        progetti = [p for p in progetti if p["name"] in filtro_progetti]
    soglia = dal - timedelta(days=MARGINE_INATTIVI)
    da_leggere, saltati = [], []
    for p in progetti:
        ua = ultima_attivita(p)
        (saltati if (not tutti and ua and ua < soglia) else da_leggere).append(p)
    log(f"Progetti: {len(progetti)} ({len(da_leggere)} da leggere, {len(saltati)} inattivi saltati)"
        f" | dal {dal:%d/%m/%Y} al {al:%d/%m/%Y}")

    agg = defaultdict(lambda: {k: 0 for k in METRICHE})
    avvisi, visti, errori = [], set(), 0
    scartate = defaultdict(int)
    orari = {}                         # chiave -> [prima, ultima] foto
    ultime_sync = {}                   # chiave -> ultimo delta sincronizzato
    geom_layer = {}                    # (progetto, layer) -> punto | linea | poligono
    foto_viste = set()                 # (progetto, nome file): ogni foto conta una volta sola
    foto_senza_ora = defaultdict(list)
    if avanzamento:
        avanzamento(0, len(da_leggere), "")

    with ThreadPoolExecutor(WORKERS) as ex:
        futuri = {ex.submit(client.delta, p["id"]): p for p in da_leggere}
        for i, fut in enumerate(as_completed(futuri), 1):
            p = futuri[fut]
            nome = p.get("name") or p["id"]
            try:
                deltas = fut.result()
            except Exception as e:
                errori += 1
                msg = f"{nome}: impossibile leggere i delta ({e})"
                avvisi.append(msg)
                log("  ! " + msg)
                if avanzamento:
                    avanzamento(i, len(da_leggere), nome)
                continue
            for d in deltas:
                did = d.get("id")
                if did:
                    if did in visti:
                        continue
                    visti.add(did)
                c = d.get("content") or {}
                layer = c.get("sourceLayerId") or c.get("localLayerId") or "?"
                g_new = (c.get("new") or {}).get("geometry")
                g_old = (c.get("old") or {}).get("geometry")
                # tipo di geometria del layer: da qualsiasi delta, anche fuori periodo
                if (nome, layer) not in geom_layer:
                    tg = tipo_geom(g_new) or tipo_geom(g_old)
                    if tg:
                        geom_layer[(nome, layer)] = tg
                quando = ts(d.get("created_at"))
                if not quando:
                    continue
                giorno = locale(quando).date()
                if not (dal <= giorno <= al):
                    continue
                stato = (d.get("last_status") or d.get("status") or "").lower()
                metodo = (c.get("method") or "").lower()
                op = d.get("created_by") or "?"
                k = (giorno, op, nome, layer)
                a = agg[k]
                ultime_sync[k] = max(ultime_sync.get(k, quando), quando)
                # foto: vanno nel giorno di scatto, a prescindere dallo stato del delta
                for f in foto_delta(c):
                    if (nome, nome_file(f)) in foto_viste:
                        continue
                    foto_viste.add((nome, nome_file(f)))
                    scatto = ora_foto(f, locale(quando))
                    if not scatto:
                        foto_senza_ora[nome].append(nome_file(f))
                        continue
                    if not (dal <= scatto.date() <= al):
                        continue
                    kf = (scatto.date(), op, nome, layer)
                    agg[kf]["foto"] += 1
                    o = orari.setdefault(kf, [scatto, scatto])
                    o[0], o[1] = min(o[0], scatto), max(o[1], scatto)
                if stato in STATI_PROBLEMA:
                    a["problemi"] += 1
                if stato not in STATI_VALIDI:
                    continue
                if metodo == "create":
                    km, mq, sc = misura(g_new)
                    a["km_netti"] += km
                    a["km_nuovi"] += km
                    a["mq_netti"] += mq
                    a["create"] += 1
                    a["punti"] += tipo_geom(g_new) == "punto"
                    scartate[nome] += sc
                elif metodo == "patch":
                    a["modificate"] += 1
                    if g_new and g_old:
                        km1, mq1, sc1 = misura(g_new)
                        km0, mq0, sc0 = misura(g_old)
                        if not (sc1 or sc0):
                            a["km_netti"] += km1 - km0
                            a["mq_netti"] += mq1 - mq0
                        scartate[nome] += sc1 or sc0
                elif metodo == "delete":
                    km, mq, sc = misura(g_old)
                    a["km_netti"] -= km
                    a["mq_netti"] -= mq
                    a["cancellate"] += 1
                    scartate[nome] += sc
            if avanzamento:
                avanzamento(i, len(da_leggere), nome)

    for nome, n in sorted(scartate.items()):
        if n:
            avvisi.append(f"{nome}: {n} geometrie scartate perché più lunghe di {KM_MAX_FEATURE} km "
                          f"(probabile CRS errato)")
    for nome, nomi in sorted(foto_senza_ora.items()):
        avvisi.append(f"{nome}: {len(nomi)} foto senza data e ora leggibili nel nome del file o con orario "
                      f"non plausibile (es. {', '.join(nomi[:3])}): non contano in foto, inizio e fine")
    if saltati:
        avvisi.append(f"{len(saltati)} progetti saltati perché senza attività dal "
                      f"{soglia:%d/%m/%Y} (usa 'Includi progetti inattivi' per leggerli tutti)")

    righe = []
    for k, v in agg.items():
        if not any(v.values()):
            continue
        g, op, pr, ly = k
        o, us = orari.get(k), ultime_sync.get(k)
        righe.append(dict(giorno=g.isoformat(), operatore=op, progetto=pr, layer=ly,
                          geom=geom_layer.get((pr, ly), ""),
                          km_netti=round(v["km_netti"], 4), km_nuovi=round(v["km_nuovi"], 4),
                          mq_netti=round(v["mq_netti"], 1), create=v["create"], punti=v["punti"],
                          modificate=v["modificate"], cancellate=v["cancellate"], problemi=v["problemi"],
                          foto=v["foto"], inizio=o and f"{o[0]:%H:%M}", fine=o and f"{o[1]:%H:%M}",
                          sync=us and f"{locale(us):%H:%M}"))
    righe.sort(key=lambda r: (r["giorno"], r["operatore"], r["progetto"], r["layer"]), reverse=True)
    return dict(righe=righe, avvisi=avvisi,
                progetti=dict(letti=len(da_leggere) - errori, saltati=len(saltati), errori=errori))


def raggruppa(righe, chiavi):
    """Somma le metriche per le chiavi indicate; aggiunge giorni_attivi, km_al_giorno,
    inizio (prima foto) e fine (ultima foto)."""
    acc = {}
    for r in righe:
        k = tuple(r[c] for c in chiavi)
        a = acc.setdefault(k, dict({c: r[c] for c in chiavi}, **{m: 0 for m in METRICHE},
                                   inizio=None, fine=None, _g=set()))
        for m in METRICHE:
            a[m] += r[m]
        if r.get("inizio"):
            a["inizio"] = min(a["inizio"] or r["inizio"], r["inizio"])
            a["fine"] = max(a["fine"] or r["fine"], r["fine"])
        if ha_sync(r):
            a["_g"].add(r["giorno"])
    out = []
    for a in acc.values():
        a["giorni_attivi"] = len(a.pop("_g"))
        a["km_al_giorno"] = a["km_netti"] / a["giorni_attivi"] if a["giorni_attivi"] else 0
        a["ore"] = ore_lavoro(a)
        out.append(a)
    return out


# ------------------------------------------------------------------ Excel ---
FORMATI = {"km": "#,##0.000", "mq": "#,##0", "int": "#,##0", "dec": "#,##0.00", "data": "dd/mm/yyyy", "ora": "hh:mm",
           "ore": "[h]:mm"}
NUMERICI = {"km", "mq", "int", "dec"}
SOMMABILI = {"km", "mq", "int", "ore"}
NON_SOMMARE = {"giorni_attivi", "n_operatori", "n_progetti"}  # conteggi distinti: la somma non ha senso


def foglio(wb, titolo, colonne, righe, totali=True):
    """
    colonne: lista di (chiave, etichetta, tipo) con tipo in testo|data|ora|ore|km|mq|int|dec.
    Header fisso, filtri automatici, riga totali con SUBTOTALE (segue i filtri di Excel).
    """
    ws = wb.create_sheet(titolo[:31])
    blu = PatternFill("solid", fgColor="1F4E79")
    ws.append([c[1] for c in colonne])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = blu
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in righe:
        riga = []
        for k, _, t in colonne:
            v = r.get(k)
            if t == "data" and isinstance(v, str):
                try:
                    v = date.fromisoformat(v)
                except ValueError:
                    pass
            elif t == "ora" and isinstance(v, str):
                v = datetime.strptime(v, "%H:%M").time()
            elif t == "ore" and v is not None:
                v = v / 24  # in Excel le durate sono frazioni di giorno
            riga.append(v)
        ws.append(riga)
    n = len(righe)
    for j, (k, etich, t) in enumerate(colonne, 1):
        L = get_column_letter(j)
        if t in FORMATI:
            for row in ws.iter_rows(min_row=2, max_row=n + 1, min_col=j, max_col=j):
                row[0].number_format = FORMATI[t]
        larg = max([len(etich)] + [len(str(r.get(k, ""))) for r in righe[:500]])
        ws.column_dimensions[L].width = min(max(11, larg + 3), 45)
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colonne))}{n + 1}"
        if totali:
            tr = n + 2
            ws.cell(tr, 1, "Totale (righe visibili)").font = Font(bold=True)
            for j, (k, _, t) in enumerate(colonne, 1):
                if t in SOMMABILI and j > 1 and k not in NON_SOMMARE:
                    L = get_column_letter(j)
                    c = ws.cell(tr, j, f"=SUBTOTAL(109,{L}2:{L}{n + 1})")
                    c.font = Font(bold=True)
                    c.number_format = FORMATI[t]
    ws.freeze_panes = "A2"
    return ws


def foglio_info(wb, voci):
    ws = wb.create_sheet("Info")
    for k, v in voci:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90
    return ws


C_GIORNO = ("giorno", "Giorno", "data")
C_OP = ("operatore", "Operatore", "testo")
C_PR = ("progetto", "Progetto", "testo")
C_LY = ("layer", "Layer", "testo")
C_TIPO = ("tipo", "Tipo", "testo")
C_MET = [("km_netti", "Km netti", "km"), ("km_nuovi", "Km nuovi", "km"), ("create", "Create", "int"),
         ("punti", "di cui punti", "int"), ("modificate", "Modificate", "int"), ("cancellate", "Cancellate", "int")]
C_GG = [("giorni_attivi", "Giorni attivi", "int"), ("km_al_giorno", "Km/giorno", "dec")]
C_ORARI = [("inizio", "Inizio (1a foto)", "ora"), ("fine", "Fine (ultima foto)", "ora"), ("ore", "Ore lavoro", "ore")]
C_FOTO = [("foto", "Foto", "int")]


def scrivi_excel(ris, dal, al, out):
    righe = ris["righe"]
    wb = Workbook()
    wb.remove(wb.active)
    ord_km = lambda xs: sorted(xs, key=lambda r: -r["km_netti"])
    foglio(wb, "Riepilogo operatori", [C_OP] + C_MET + C_FOTO + C_GG, ord_km(raggruppa(righe, ["operatore"])))
    foglio(wb, "Operatore x giorno", [C_GIORNO, C_OP] + C_ORARI + C_MET + C_FOTO,
           sorted(raggruppa(righe, ["giorno", "operatore"]), key=lambda r: (r["giorno"], r["operatore"])))
    foglio(wb, "Progetto x operatore", [C_PR, C_OP] + C_MET + C_FOTO + C_GG,
           sorted(raggruppa(righe, ["progetto", "operatore"]), key=lambda r: (r["progetto"], -r["km_netti"])))
    foglio(wb, "Progetto x layer", [C_PR, C_LY, C_TIPO] + C_MET + C_FOTO,
           sorted((dict(r, tipo=TIPI[r["geom"]]) for r in raggruppa(righe, ["progetto", "layer", "geom"])),
                  key=lambda r: (r["progetto"], -r["km_netti"])))
    foglio(wb, "Dettaglio", [C_GIORNO, C_OP, C_PR, C_LY, C_TIPO] + C_ORARI + C_MET + C_FOTO,
           sorted((dict(r, ore=ore_lavoro(r), tipo=TIPI[r["geom"]]) for r in righe),
                  key=lambda r: (r["giorno"], r["operatore"], r["progetto"], r["layer"])))
    foglio_info(wb, [
        ("Periodo", f"{dal:%d/%m/%Y} - {al:%d/%m/%Y} (giorno di sincronizzazione, ora italiana)"),
        ("Generato il", f"{adesso():%d/%m/%Y %H:%M}"),
        ("Km netti", "linee nuove + allungamenti - accorciamenti - cancellazioni"),
        ("Km nuovi", "solo linee create nel periodo (lordo)"),
        ("di cui punti", "elementi puntuali creati (es. nodi, pali, pozzetti, note), già compresi in Create"),
        ("Tipo", "geometria del layer (punto, linea, poligono); n.d. = tabella senza geometria o nessuna "
                 "geometria nei dati"),
        ("Inizio / Fine", "orario della prima e dell'ultima foto della giornata, letto dal nome del file "
                          "(orologio del dispositivo); foto, inizio e fine sono nel giorno di SCATTO"),
        ("Ore lavoro", "Fine - Inizio: tempo tra la prima e l'ultima foto della giornata (pause comprese)"),
        ("Progetti", f"letti {ris['progetti']['letti']}, saltati {ris['progetti']['saltati']}, "
                     f"errori {ris['progetti']['errori']}"),
    ] + [("Avviso", a) for a in ris["avvisi"]])
    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description="Report produzione rilievi QFieldCloud")
    ap.add_argument("--dal", type=date.fromisoformat, default=None, help="AAAA-MM-GG (default: 7 giorni fa)")
    ap.add_argument("--al", type=date.fromisoformat, default=None, help="AAAA-MM-GG (default: oggi)")
    ap.add_argument("--progetti", nargs="*", help="nomi progetto da includere (default: tutti)")
    ap.add_argument("--tutti", action="store_true", help="leggi anche i progetti inattivi")
    ap.add_argument("--giorni", type=int, default=None, help="ultimi N giorni fino ad --al (alternativa a --dal)")
    ap.add_argument("--out", default=None, help="file Excel di output")
    ap.add_argument("--json", default=None, help="file JSON per la board web (senza --out non crea l'Excel)")
    a = ap.parse_args()
    al = a.al or oggi()
    dal = a.dal or (al - timedelta(days=max(1, a.giorni or 7) - 1))
    t0 = time.time()
    try:
        client = QFC(*carica_env())
    except Exception as e:
        sys.exit(str(e))
    ris = elabora(client, dal, al, set(a.progetti or []), tutti=a.tutti)
    tot = sum(r["km_netti"] for r in ris["righe"])
    print(f"{len(ris['righe'])} righe, {tot:.2f} km netti")
    if a.json:
        dati = dict(ris, dal=dal.isoformat(), al=al.isoformat(), giorni=(al - dal).days + 1, tutti=a.tutti,
                    aggiornato=adesso().isoformat(timespec="seconds"), durata=round(time.time() - t0, 1))
        out_json = Path(a.json)
        out_json.parent.mkdir(parents=True, exist_ok=True)
        out_json.write_text(json.dumps(dati, ensure_ascii=False), encoding="utf-8")
        print(f"Dati board salvati: {out_json}")
    if a.out or not a.json:
        out = a.out or str(Path(__file__).with_name(f"report_rilievi_{dal}_{al}.xlsx"))
        scrivi_excel(ris, dal, al, out)
        print(f"Report salvato: {out}")


if __name__ == "__main__":
    main()
